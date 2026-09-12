"""Regression contract for the read-only live campaign configuration export."""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts/export-level-configuration.py"
spec = importlib.util.spec_from_file_location("export_level_configuration", SCRIPT)
assert spec and spec.loader
EXPORT = importlib.util.module_from_spec(spec)
spec.loader.exec_module(EXPORT)


class LevelConfigurationExportTests(unittest.TestCase):
    def test_final_mr_target_and_bonus_export_independently(self):
        levels = EXPORT.read_kv(EXPORT.RUNTIME_PATH.read_text(encoding='utf-8'))['levels']
        boss = next(iter(levels['ch20']['enemies'].values()))
        boss['boss_magic_resistance_pct'] = '80'
        boss['boss_magic_resistance_bonus_pct'] = '0'
        row = next(row for row in EXPORT.unit_rows(levels) if row['关卡'] == 'ch20')
        self.assertEqual((row['Boss最终魔抗目标%'], row['Boss魔抗乘算加成%']), (80, 0))
        row = next(row for row in EXPORT.unit_rows(levels) if row['关卡'] == 'ch30')
        self.assertEqual((row['Boss最终魔抗目标%'], row['Boss魔抗乘算加成%']), ('', 80))

    def test_live_kv_export_has_all_review_tables_and_real_values(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            files = EXPORT.export(output)
            self.assertTrue(all(path.is_file() for path in files.values()))
            manifest = json.loads(files["manifest"].read_text(encoding="utf-8"))
            self.assertEqual(manifest["runtime_source"], "game/dota_addons/dota2_rpg/scripts/data/levels.kv")
            self.assertEqual(manifest["stage_count"], 30)
            self.assertEqual(manifest["unit_configuration_rows"], 125)
            self.assertEqual(manifest["equipment_rows"], 499)
            self.assertEqual(manifest["unique_units"], 75)
            self.assertEqual(manifest["item_name_count"], 76)
            self.assertEqual(manifest["untranslated_item_ids"], [])
            # Equipment authoring synchronizes actual hero levels from KV.
            self.assertEqual(manifest["maintenance_source_difference_rows"], 0)

            book = load_workbook(files["excel"], data_only=True)
            self.assertEqual(book.sheetnames, ["说明", "关卡总览", "单位明细", "装备明细", "单位出现汇总", "源文件差异"])
            for name in book.sheetnames[1:]:
                sheet = book[name]
                if name == "源文件差异":
                    self.assertEqual(sheet.cell(1, 1).value, "无数据")
                    continue
                self.assertTrue(sheet.auto_filter.ref)
                self.assertTrue(sheet.tables, name + " must be filterable")
                self.assertEqual(sheet.freeze_panes, "A2")

            unit_sheet = book["单位明细"]
            headers = [cell.value for cell in unit_sheet[1]]
            records = [dict(zip(headers, values)) for values in unit_sheet.iter_rows(min_row=2, values_only=True)]
            axe = next(row for row in records if row["关卡"] == "ch05" and row["单位原生ID"] == "npc_dota_hero_axe")
            self.assertEqual((axe["单位名称"], axe["数量"], axe["等级"], axe["AI类型"],
                              axe["装备1（中文）"], axe["装备1（原生ID）"], axe["装备2（中文）"], axe["装备2（原生ID）"]),
                             ("斧王", 1, 8, "aggro_front", "相位鞋", "item_phase_boots", "闪烁匕首", "item_blink"))
            for chapter, unit, health, attack, spell, cooldown in [
                ("ch10", "npc_dota_hero_centaur", 6000, 16.666667, 16.666667, 4.166667),
                ("ch20", "npc_dota_hero_spirit_breaker", 20000, 75, 25, 6.666667),
                ("ch30", "npc_dota_hero_skeleton_king", 50000, 100, 33.333333, 8.333333),
            ]:
                bosses = [row for row in records if row["关卡"] == chapter]
                self.assertEqual(len(bosses), 1, chapter + " must remain a solo Boss")
                boss = bosses[0]
                self.assertEqual((boss["单位原生ID"], boss["数量"], boss["是否Boss"], boss["Boss最大生命"], boss["Boss生命倍率"], boss["Boss攻击伤害+%"], boss["Boss法术增幅+%"], boss["Boss冷却减少%"]),
                                 (unit, 1, "是", health, None, attack, spell, cooldown))
                self.assertEqual((boss['Boss额外护甲'], boss['Boss魔抗乘算加成%'], boss['Boss最终魔抗目标%']),
                                 {'ch10': (None, None, None), 'ch20': (15, 0, 80), 'ch30': (100, 80, None)}[chapter])
            self.assertEqual(boss["等级"], 30)
            for chapter, expected in {'ch11': (4.5, 2.5, 20, 55), 'ch16': (6, 3.5, 30, 62),
                                      'ch21': (8.5, 4.5, 42, 70), 'ch26': (11, 6, 56, 75)}.items():
                for row in records:
                    if row['关卡'] == chapter:
                        self.assertEqual(tuple(row[key] for key in ('生命倍率', '攻击倍率', '额外护甲', '魔法抗性%')), expected)
                        self.assertIn(row['单位原生ID'], EXPORT.UNIT_NAMES)

            difference_sheet = book["源文件差异"]
            difference_headers = [cell.value for cell in difference_sheet[1]]
            differences = [dict(zip(difference_headers, values)) for values in difference_sheet.iter_rows(min_row=2, values_only=True)]
            self.assertEqual(differences, [])
            # Reporting still detects stale JSON without changing live values.
            with tempfile.TemporaryDirectory() as stale_dir:
                stale = json.loads(EXPORT.SOURCE_PATH.read_text(encoding="utf-8"))
                stale["ch05"]["enemies"][0]["level"] = 30
                stale_path = Path(stale_dir) / "stale.json"
                stale_path.write_text(json.dumps(stale), encoding="utf-8")
                levels = EXPORT.read_kv(EXPORT.RUNTIME_PATH.read_text(encoding="utf-8"))["levels"]
                delta = EXPORT.compare_maintenance_source(levels, stale_path)
                self.assertEqual(len(delta), 1)
                self.assertEqual((delta[0]["运行时 levels.kv"], delta[0]["维护 levels_v07.json"]), ("8", "30"))

            equipment_sheet = book["装备明细"]
            equipment_headers = [cell.value for cell in equipment_sheet[1]]
            equipment = [dict(zip(equipment_headers, values)) for values in equipment_sheet.iter_rows(min_row=2, values_only=True)]
            self.assertTrue(all(row["装备中文名"] and row["装备原生ID"] for row in equipment))
            self.assertTrue({row["装备原生ID"] for row in equipment} <= set(EXPORT.ITEM_NAMES))
            self.assertEqual(len(equipment), sum(row["装备数量"] for row in records))
            self.assertTrue(any(row["装备槽"] == 6 for row in equipment))
            self.assertTrue(any(row["装备6（原生ID）"] for row in records))
            self.assertEqual(next(row["装备中文名"] for row in equipment if row["装备原生ID"] == "item_blink"), "闪烁匕首")
            self.assertEqual(next(row["装备中文名"] for row in equipment if row["装备原生ID"] == "item_heavens_halberd"), "天堂之戟")

            raw = files["units_csv"].read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "CSV needs an Excel-friendly UTF-8 BOM")
            with files["units_csv"].open(encoding="utf-8-sig", newline="") as stream:
                csv_rows = list(csv.DictReader(stream))
            self.assertEqual(len(csv_rows), 125)
            self.assertEqual(csv_rows[0]["关卡"], "ch01")
            expected_hash = hashlib.sha256(EXPORT.RUNTIME_PATH.read_bytes()).hexdigest()
            self.assertEqual(manifest["runtime_sha256"], expected_hash)


if __name__ == "__main__":
    unittest.main(verbosity=2)
