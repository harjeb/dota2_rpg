#!/usr/bin/env python3
"""Export the live campaign configuration to reviewable Excel and CSV tables.

The game loads scripts/data/levels.kv at runtime. This tool deliberately reads
that file, not the historical levels.json. It never changes campaign data.

    python scripts/export-level-configuration.py
    python scripts/export-level-configuration.py --output-dir exports/level_configuration_current
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
RUNTIME_PATH = ROOT / "game/dota_addons/dota2_rpg/scripts/data/levels.kv"
SOURCE_PATH = ROOT / "game/dota_addons/dota2_rpg/scripts/data/levels_v07.json"
DEFAULT_OUTPUT = ROOT / "exports/level_configuration_current"

# Only names present in the current runtime configuration are translated. The
# native ID stays in every row and is the authoritative edit key.
UNIT_NAMES = {
    "npc_dota_hero_abaddon": "亚巴顿", "npc_dota_hero_ancient_apparition": "远古冰魄",
    "npc_dota_hero_antimage": "敌法师", "npc_dota_hero_axe": "斧王",
    "npc_dota_hero_bane": "祸乱之源", "npc_dota_hero_bloodseeker": "血魔",
    "npc_dota_hero_bristleback": "钢背兽", "npc_dota_hero_centaur": "半人马战行者",
    "npc_dota_hero_chaos_knight": "混沌骑士", "npc_dota_hero_clinkz": "克林克兹",
    "npc_dota_hero_crystal_maiden": "水晶室女", "npc_dota_hero_dark_willow": "邪影芳灵",
    "npc_dota_hero_dazzle": "戴泽", "npc_dota_hero_death_prophet": "死亡先知",
    "npc_dota_hero_disruptor": "干扰者", "npc_dota_hero_dragon_knight": "龙骑士",
    "npc_dota_hero_drow_ranger": "卓尔游侠", "npc_dota_hero_enchantress": "魅惑魔女",
    "npc_dota_hero_grimstroke": "天涯墨客", "npc_dota_hero_gyrocopter": "矮人直升机",
    "npc_dota_hero_huskar": "哈斯卡", "npc_dota_hero_jakiro": "杰奇洛",
    "npc_dota_hero_juggernaut": "主宰", "npc_dota_hero_leshrac": "拉席克",
    "npc_dota_hero_lich": "巫妖", "npc_dota_hero_life_stealer": "噬魂鬼",
    "npc_dota_hero_lina": "莉娜", "npc_dota_hero_lion": "莱恩",
    "npc_dota_hero_luna": "露娜", "npc_dota_hero_necrolyte": "瘟疫法师",
    "npc_dota_hero_nevermore": "影魔", "npc_dota_hero_night_stalker": "暗夜魔王",
    "npc_dota_hero_ogre_magi": "食人魔魔法师", "npc_dota_hero_omniknight": "全能骑士",
    "npc_dota_hero_oracle": "神谕者", "npc_dota_hero_phantom_assassin": "幻影刺客",
    "npc_dota_hero_phantom_lancer": "幻影长矛手", "npc_dota_hero_pugna": "帕格纳",
    "npc_dota_hero_queenofpain": "痛苦女王", "npc_dota_hero_razor": "剃刀",
    "npc_dota_hero_riki": "力丸", "npc_dota_hero_sand_king": "沙王",
    "npc_dota_hero_shadow_demon": "暗影恶魔", "npc_dota_hero_shadow_shaman": "暗影萨满",
    "npc_dota_hero_silencer": "沉默术士", "npc_dota_hero_skeleton_king": "冥魂大帝",
    "npc_dota_hero_skywrath_mage": "天怒法师", "npc_dota_hero_slardar": "斯拉达",
    "npc_dota_hero_slark": "斯拉克", "npc_dota_hero_sniper": "狙击手",
    "npc_dota_hero_spirit_breaker": "裂魂人", "npc_dota_hero_sven": "斯温",
    "npc_dota_hero_templar_assassin": "圣堂刺客", "npc_dota_hero_tidehunter": "潮汐猎人",
    "npc_dota_hero_treant": "树精卫士", "npc_dota_hero_troll_warlord": "巨魔战将",
    "npc_dota_hero_undying": "不朽尸王", "npc_dota_hero_ursa": "熊战士",
    "npc_dota_hero_vengefulspirit": "复仇之魂", "npc_dota_hero_venomancer": "剧毒术士",
    "npc_dota_hero_viper": "冥界亚龙", "npc_dota_hero_warlock": "术士",
    "npc_dota_hero_witch_doctor": "巫医", "npc_dota_hero_zuus": "宙斯",
    "npc_dota_neutral_alpha_wolf": "头狼", "npc_dota_neutral_black_dragon": "黑龙",
    "npc_dota_neutral_centaur_khan": "半人马可汗", "npc_dota_neutral_dark_troll_warlord": "黑暗巨魔首领",
    "npc_dota_neutral_gnoll_assassin": "豺狼刺客", "npc_dota_neutral_kobold": "狗头人",
    "npc_dota_neutral_ogre_mauler": "食人魔拳手", "npc_dota_neutral_polar_furbolg_champion": "极地熊怪勇士",
    "npc_dota_neutral_satyr_hellcaller": "萨特地狱使者",
}
AI_NAMES = {
    "simple_nearest": "最近目标（野怪）", "aggro_front": "前排近距攻击",
    "focus_lowest_hp": "优先最低生命", "ai_healer_protect": "治疗/保护友军",
}
# Snapshot of the current native abilities_schinese localization for every item
# actually referenced by levels.kv. Keep IDs too: the KV accepts IDs, not labels.
ITEM_NAMES = {
    "item_abyssal_blade": "深渊之刃", "item_aeon_disk": "永恒之盘", "item_aether_lens": "以太透镜",
    "item_arcane_boots": "奥术鞋", "item_armlet": "莫尔迪基安的臂章", "item_assault": "强袭胸甲",
    "item_basher": "碎颅锤", "item_bfury": "狂战斧", "item_black_king_bar": "黑皇杖",
    "item_blade_mail": "刃甲", "item_blink": "闪烁匕首", "item_bloodstone": "血精石",
    "item_bloodthorn": "血棘", "item_boots": "速度之靴", "item_bracer": "护腕",
    "item_butterfly": "蝴蝶", "item_crimson_guard": "赤红甲", "item_cyclone": "Eul的神圣法杖",
    "item_desolator": "黯灭", "item_diffusal_blade": "净魂之刃", "item_disperser": "散魂剑",
    "item_dragon_lance": "魔龙枪", "item_echo_sabre": "回音战刃", "item_eternal_shroud": "永世法衣",
    "item_force_staff": "原力法杖", "item_glimmer_cape": "微光披风", "item_greater_crit": "代达罗斯之殇",
    "item_guardian_greaves": "卫士胫甲", "item_halberd": "天堂之戟（兼容旧 ID）", "item_harpoon": "鱼叉",
    "item_heart": "恐鳌之心", "item_hurricane_pike": "飓风长戟", "item_invis_sword": "影刃",
    "item_kaya_and_sange": "散慧对剑", "item_lotus_orb": "清莲宝珠", "item_maelstrom": "漩涡",
    "item_magic_wand": "魔杖", "item_manta": "幻影斧", "item_mekansm": "梅肯斯姆",
    "item_meteor_hammer": "陨星锤", "item_mjollnir": "雷神之锤", "item_null_talisman": "空灵挂件",
    "item_octarine_core": "玲珑心", "item_orchid": "紫怨", "item_phase_boots": "相位鞋",
    "item_phylactery": "灵匣", "item_pipe": "洞察烟斗", "item_power_treads": "动力鞋",
    "item_refresher": "刷新球", "item_rod_of_atos": "阿托斯之棍", "item_sange_and_yasha": "散夜对剑",
    "item_satanic": "撒旦之邪力", "item_sheepstick": "邪恶镰刀", "item_shivas_guard": "希瓦的守护",
    "item_skadi": "斯嘉蒂之眼", "item_spirit_vessel": "魂之灵瓮", "item_ultimate_scepter": "阿哈利姆神杖",
    "item_wind_lace": "风灵之纹", "item_wind_waker": "风之杖", "item_wraith_band": "怨灵系带",
    "item_yasha": "夜叉",
}


def read_kv(text: str) -> dict[str, Any]:
    """Parse the deliberately simple quoted KV used by levels.kv."""
    import shlex
    tokens = iter(shlex.split(text))

    def block() -> dict[str, Any]:
        result: dict[str, Any] = {}
        for key in tokens:
            if key == "}":
                return result
            value = next(tokens)
            result[key] = block() if value == "{" else value
        return result
    return block()


def chapter_number(stage_id: str) -> int:
    match = re.fullmatch(r"ch(\d+)", stage_id)
    if not match:
        raise ValueError(f"Unexpected chapter ID: {stage_id}")
    return int(match.group(1))


def ordered_values(value: Any) -> list[Any]:
    if not isinstance(value, dict):
        return []
    return [value[key] for key in sorted(value, key=lambda key: int(key) if str(key).isdigit() else str(key))]


def as_number(value: Any) -> int | float | str:
    text = str(value)
    try:
        number = float(text)
    except ValueError:
        return text
    return int(number) if number.is_integer() else number


def unit_kind(unit: str) -> str:
    return "英雄" if unit.startswith("npc_dota_hero_") else "野怪"


def native_short(unit: str) -> str:
    return re.sub(r"^npc_dota_(?:hero|neutral)_", "", unit)


def row_items(enemy: dict[str, Any]) -> list[str]:
    return [str(item) for item in ordered_values(enemy.get("items", {}))]


def item_name(item_id: str) -> str:
    return ITEM_NAMES.get(item_id, "未收录中文名（" + item_id + "）")


def row_tags(enemy: dict[str, Any]) -> list[str]:
    return [str(tag) for tag in ordered_values(enemy.get("tags", {}))]


def stage_rows(levels: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for stage_id in sorted(levels, key=chapter_number):
        stage = levels[stage_id]
        enemies = ordered_values(stage["enemies"])
        reward = stage.get("reward", {})
        rows.append({
            "关卡": stage_id, "关卡序号": chapter_number(stage_id), "关卡名称": stage.get("name", ""),
            "类型": stage.get("type", ""), "推荐等级": as_number(stage.get("recommended_level", "")),
            "关卡倍率": as_number(stage.get("multi", "")), "敌方配置行数": len(enemies),
            "敌方总数量": sum(int(enemy.get("count", 1)) for enemy in enemies),
            "英雄配置行数": sum(unit_kind(str(enemy["unit"])) == "英雄" for enemy in enemies),
            "野怪配置行数": sum(unit_kind(str(enemy["unit"])) == "野怪" for enemy in enemies),
            "首通金币": as_number(reward.get("gold", "")),
            "每上阵英雄经验": as_number(reward.get("xp_per_active_hero", "")),
            "限时秒数": as_number(stage.get("time_limit", "")),
            "时间奖励上限": as_number(stage.get("time_bonus_cap", "")),
            "战利品表": stage.get("loot", ""), "配置路径": f"levels/{stage_id}",
        })
    return rows


def unit_rows(levels: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for stage_id in sorted(levels, key=chapter_number):
        stage = levels[stage_id]
        for config_index, enemy in enumerate(ordered_values(stage["enemies"]), 1):
            unit = str(enemy["unit"])
            items, tags = row_items(enemy), row_tags(enemy)
            row = {
                "关卡": stage_id, "关卡序号": chapter_number(stage_id), "关卡名称": stage.get("name", ""),
                "配置序号": config_index, "单位类型": unit_kind(unit), "单位名称": UNIT_NAMES.get(unit, native_short(unit)),
                "单位原生ID": unit, "数量": as_number(enemy.get("count", 1)), "等级": as_number(enemy.get("level", "")),
                "AI类型": enemy.get("ai", ""), "AI说明": AI_NAMES.get(str(enemy.get("ai", "")), "未列入说明"),
                "标签": "、".join(tags), "是否Boss": "是" if "boss" in tags else "否",
                "模板": enemy.get("template", ""), "生命倍率": as_number(enemy.get("hp_multiplier", "")),
                "攻击倍率": as_number(enemy.get("attack_multiplier", "")), "额外护甲": as_number(enemy.get("bonus_armor", "")),
                "魔法抗性%": as_number(enemy.get("magic_resistance", "")),
                "状态抗性%": as_number(enemy.get("status_resistance", "")),
                "Boss最大生命": as_number(enemy.get("boss_max_health", "")),
                "Boss生命倍率": as_number(enemy.get("boss_health_multiplier", "")),
                "Boss攻击伤害+%": as_number(enemy.get("boss_attack_damage_pct", "")),
                "Boss法术增幅+%": as_number(enemy.get("boss_spell_amp_pct", "")),
                "Boss冷却减少%": as_number(enemy.get("boss_cooldown_reduction_pct", "")),
                "装备数量": len(items), "装备合计（中文）": "；".join(item_name(item) for item in items),
                "装备合计（原生ID）": "；".join(items), "配置路径": f"levels/{stage_id}/enemies/{config_index}",
            }
            for item_index in range(1, 6):
                item = items[item_index - 1] if item_index <= len(items) else ""
                row[f"装备{item_index}（中文）"] = item_name(item) if item else ""
                row[f"装备{item_index}（原生ID）"] = item
            rows.append(row)
    return rows


def equipment_rows(units: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for unit in units:
        for item_index in range(1, 6):
            item = unit[f"装备{item_index}（原生ID）"]
            if item:
                rows.append({
                    "关卡": unit["关卡"], "关卡序号": unit["关卡序号"], "配置序号": unit["配置序号"],
                    "单位名称": unit["单位名称"], "单位原生ID": unit["单位原生ID"], "单位等级": unit["等级"],
                    "标签": unit["标签"], "装备槽": item_index, "装备中文名": item_name(item), "装备原生ID": item,
                    "配置路径": unit["配置路径"] + f"/items/{item_index}",
                })
    return rows


def appearance_rows(units: list[dict[str, Any]]) -> list[dict[str, Any]]:
    appearances: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in units:
        appearances[row["单位原生ID"]].append(row)
    rows = []
    for unit, entries in sorted(appearances.items(), key=lambda pair: (pair[1][0]["单位类型"], pair[1][0]["单位名称"])):
        rows.append({
            "单位类型": entries[0]["单位类型"], "单位名称": entries[0]["单位名称"], "单位原生ID": unit,
            "出现配置次数": len(entries), "累计刷出数量": sum(int(row["数量"]) for row in entries),
            "出现关卡": "、".join(row["关卡"] for row in entries),
            "等级范围": f"{min(float(row['等级']) for row in entries):g}–{max(float(row['等级']) for row in entries):g}",
            "Boss关卡": "、".join(row["关卡"] for row in entries if row["是否Boss"] == "是"),
        })
    return rows


def normalize(value: Any) -> Any:
    if isinstance(value, dict):
        if value and all(str(key).isdigit() for key in value):
            return [normalize(value[key]) for key in sorted(value, key=lambda key: int(key))]
        return {key: normalize(item) for key, item in value.items()}
    if isinstance(value, list):
        return [normalize(item) for item in value]
    return str(value)


def comparison_value(value: Any) -> Any:
    normalized = normalize(value)
    return normalized if isinstance(normalized, str) else json.dumps(normalized, ensure_ascii=False)


def compare_maintenance_source(levels: dict[str, Any], source_path: Path) -> list[dict[str, Any]]:
    """Report, but never substitute, the optional historical JSON source.

    `levels.kv` is what DataLoader reads. `levels_v07.json` remains useful to
    author roster/equipment changes, but it can lag the live KV. Making this a
    visible sheet avoids exporting a misleading hybrid or silently overwriting
    the runtime values a user asked to review.
    """
    source = json.loads(source_path.read_text(encoding="utf-8"))
    differences: list[dict[str, Any]] = []
    for stage_id in sorted(set(levels) | set(source), key=chapter_number):
        if stage_id not in levels or stage_id not in source:
            differences.append({"关卡": stage_id, "配置序号": "", "单位原生ID": "", "字段": "关卡存在性",
                                "运行时 levels.kv": "存在" if stage_id in levels else "缺少",
                                "维护 levels_v07.json": "存在" if stage_id in source else "缺少"})
            continue
        runtime_enemies = ordered_values(levels[stage_id]["enemies"])
        source_enemies = source[stage_id].get("enemies", [])
        if len(runtime_enemies) != len(source_enemies):
            differences.append({"关卡": stage_id, "配置序号": "", "单位原生ID": "", "字段": "敌方配置行数",
                                "运行时 levels.kv": len(runtime_enemies), "维护 levels_v07.json": len(source_enemies)})
        for index, (runtime, authored) in enumerate(zip(runtime_enemies, source_enemies), 1):
            unit = str(runtime.get("unit", authored.get("unit", "")))
            for field in sorted(set(runtime) | set(authored)):
                actual, expected = runtime.get(field, "（缺少）"), authored.get(field, "（缺少）")
                if normalize(actual) != normalize(expected):
                    differences.append({"关卡": stage_id, "配置序号": index, "单位原生ID": unit, "字段": field,
                                        "运行时 levels.kv": comparison_value(actual),
                                        "维护 levels_v07.json": comparison_value(expected)})
    return differences


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path: Path, sheets: list[tuple[str, list[dict[str, Any]]]], metadata: dict[str, str]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise RuntimeError("Excel export requires openpyxl; install it with: python -m pip install openpyxl") from error

    book = Workbook()
    book.remove(book.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    boss_fill = PatternFill("solid", fgColor="FCE4D6")
    for index, (sheet_name, rows) in enumerate(sheets):
        sheet = book.create_sheet(sheet_name)
        if not rows:
            sheet.append(["无数据"])
            continue
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        table = Table(displayName=f"LevelConfigTable{index + 1}", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
        for column, header in enumerate(headers, 1):
            width = max(len(str(header)), *(len(str(row.get(header, ""))) for row in rows))
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 11), 44)
        sheet.row_dimensions[1].height = 33
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if "是否Boss" in headers:
            boss_column = headers.index("是否Boss") + 1
            for row_index in range(2, len(rows) + 2):
                if sheet.cell(row_index, boss_column).value == "是":
                    for cell in sheet[row_index]: cell.fill = boss_fill

    info = book.create_sheet("说明", 0)
    info.append(["当前关卡配置导出（只读快照）"])
    info["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    info["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    info.merge_cells("A1:D1")
    notes = [
        ("运行时来源", metadata["runtime_path"]),
        ("来源 SHA256", metadata["runtime_sha256"]),
        ("运行时优先", "游戏实际读取 levels.kv。本表绝不使用 levels_v07.json 覆盖运行时数值。若存在差异，请看“源文件差异”表。"),
        ("如何手动调整", "不要编辑本 Excel/CSV 让游戏生效；请编辑 levels.kv，并按需要同步更新 levels_v07.json，再运行本导出工具复查。"),
        ("关卡总览", "每关一行：奖励、推荐等级、限时、总数量、倍率与战利品表。"),
        ("装备中文名", "中文名来自当前 Dota 简体中文物品本地化快照；“原生ID”仍是手动改 KV 时必须使用的名称。item_halberd 是兼容旧 ID，特别标注。"),
        ("单位明细", "每关每个 enemies 配置一行。配置路径直接对应 levels.kv 的 enemies/N。装备1–5按原生槽位导出中文名及原生ID。"),
        ("装备明细", "每件装备一行，便于按中文名、原生ID或关卡筛选。"),
        ("单位出现汇总", "按单位汇总当前所有关卡的配置次数、实际累计刷出数量、等级范围与 Boss 出现关卡。"),
        ("AI说明", "simple_nearest=野怪最近目标；aggro_front=前排近距攻击；focus_lowest_hp=优先最低生命；ai_healer_protect=治疗/保护友军。"),
        ("Boss列", "仅 Boss 行有 Boss最大生命（优先于生命倍率）、攻击伤害、法术增幅、冷却减少数值；这些由 modifier_rpg_boss_power 生效。"),
    ]
    for key, value in notes:
        info.append([key, value])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 112
    for row in range(2, len(notes) + 2):
        info.cell(row, 1).font = Font(bold=True)
        info.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        info.row_dimensions[row].height = 32
    info.freeze_panes = "A2"
    book.properties.title = "当前关卡单位、等级与装备配置"
    book.properties.description = "从运行时 levels.kv 导出的只读平衡审阅表"
    book.save(path)


def export(output_dir: Path) -> dict[str, Path]:
    levels = read_kv(RUNTIME_PATH.read_text(encoding="utf-8"))["levels"]
    source_differences = compare_maintenance_source(levels, SOURCE_PATH)
    stages = stage_rows(levels)
    units = unit_rows(levels)
    equipment = equipment_rows(units)
    appearances = appearance_rows(units)
    output_dir.mkdir(parents=True, exist_ok=True)
    source_hash = hashlib.sha256(RUNTIME_PATH.read_bytes()).hexdigest()
    outputs = {
        "excel": output_dir / "当前关卡配置.xlsx",
        "stages_csv": output_dir / "关卡总览.csv",
        "units_csv": output_dir / "单位明细.csv",
        "equipment_csv": output_dir / "装备明细.csv",
        "appearance_csv": output_dir / "单位出现汇总.csv",
        "source_differences_csv": output_dir / "源文件差异.csv",
    }
    write_workbook(outputs["excel"], [("关卡总览", stages), ("单位明细", units), ("装备明细", equipment), ("单位出现汇总", appearances), ("源文件差异", source_differences)], {
        "runtime_path": str(RUNTIME_PATH.relative_to(ROOT)).replace("\\", "/"), "runtime_sha256": source_hash,
    })
    for key, rows in (("stages_csv", stages), ("units_csv", units), ("equipment_csv", equipment), ("appearance_csv", appearances), ("source_differences_csv", source_differences)):
        write_csv(outputs[key], rows)
    manifest = {
        "runtime_source": str(RUNTIME_PATH.relative_to(ROOT)).replace("\\", "/"),
        "runtime_sha256": source_hash,
        "validated_against": str(SOURCE_PATH.relative_to(ROOT)).replace("\\", "/"),
        "stage_count": len(stages), "unit_configuration_rows": len(units),
        "equipment_rows": len(equipment), "unique_units": len(appearances),
        "item_name_count": len(ITEM_NAMES),
        "untranslated_item_ids": sorted({row["装备原生ID"] for row in equipment if row["装备原生ID"] not in ITEM_NAMES}),
        "maintenance_source_difference_rows": len(source_differences),
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "files": {key: path.name for key, path in outputs.items()},
    }
    manifest_path = output_dir / "导出清单.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    outputs["manifest"] = manifest_path
    return outputs


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT, help="output directory (default: exports/level_configuration_current)")
    args = parser.parse_args()
    outputs = export(args.output_dir.resolve())
    print("Exported runtime campaign configuration:")
    for key, path in outputs.items(): print(f"  {key}: {path}")


if __name__ == "__main__":
    main()
